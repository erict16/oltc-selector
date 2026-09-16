#!/usr/bin/env python3
"""Write 2025/2026 OS Excel workbooks with replay duty columns.

Reads docs/replay/YYYY-os-sales.json (from OS PDF extract). Does not invent rows.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
REPLAY = ROOT / "docs" / "replay"

DUTY_COLS = [
    "file",
    "folder",
    "serial",
    "customer",
    "family_folder",
    "sold_type",
    "type_alts",
    "phases",
    "current_a",
    "connection",
    "um_kv",
    "selector_size",
    "tap_code",
    "unit_count",
    "mva",
    "rated_kv",
    "i_a",
    "i_max_a",
    "ust_v",
    "ust_max_v",
    "plus_minus_steps",
    "plus_steps",
    "minus_steps",
    "mdu",
    "text_status",
    "notes",
]


def load_sales(year: int) -> list[dict]:
    path = REPLAY / f"{year}-os-sales.json"
    if not path.exists():
        raise SystemExit(f"missing {path} — extract OS first, do not invent rows")
    raw = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(raw, list):
        return raw
    if isinstance(raw, dict) and isinstance(raw.get("sales"), list):
        return raw["sales"]
    raise SystemExit(f"unrecognised JSON in {path}")


def write_xlsx(year: int, rows: list[dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "sales"
    ws.append(DUTY_COLS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for rec in rows:
        ws.append([rec.get(c) for c in DUTY_COLS])
    for i, col in enumerate(DUTY_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(col) + 2, 12), 28)
    summary = wb.create_sheet("summary")
    summary.append(["year", year])
    summary.append(["rows", len(rows)])
    summary.append(["with_sold_type", sum(1 for r in rows if r.get("sold_type"))])
    summary.append(["with_i_max", sum(1 for r in rows if r.get("i_max_a") or r.get("i_a"))])
    out = REPLAY / f"{year}-os-sales.xlsx"
    wb.save(out)
    return out


def main() -> None:
    for year in (2025, 2026):
        rows = load_sales(year)
        if not rows:
            raise SystemExit(f"{year} OS JSON is empty")
        path = write_xlsx(year, rows)
        print(f"wrote {path} rows={len(rows)}")


if __name__ == "__main__":
    main()
