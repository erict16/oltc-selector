#!/usr/bin/env python3
"""Build lib/listIndex.data.json from the 2025 list: WSL/WDL keys, no RMB.

Never commit the xlsx. A CMA7 cell must be present so the row is a real
list line; the number itself is discarded.
"""
from __future__ import annotations

import json
import re
import sys
from datetime import date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("need openpyxl")

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "lib" / "listIndex.data.json"
ONEDRIVE = (
    Path.home()
    / "Library/CloudStorage/OneDrive-上海华明电力设备制造有限公司"
    / "QS"
    / "a. Base Price List 2025.xlsx"
)

OCTC_ROMAN = r"(?P<series>VIII|VII|III|II|IV|VI|IX|V|I)"
OCTC_UMS = (
    "363",
    "362",
    "330",
    "300",
    "252",
    "170",
    "145",
    "126",
    "72.5",
    "40.5",
    "17.5",
    "12.5",
    "12",
)


def compact_octc_label(raw: str) -> str:
    s = raw.replace("\xa0", " ")
    s = re.sub(r"\s+", "", s)
    s = s.replace("×", "x").replace("*", "x").replace("X", "x")
    s = s.replace("（", "(").replace("）", ")")
    s = s.replace("－", "-").replace("–", "-").replace("—", "-")
    s = s.replace("／", "/")
    s = s.replace(",", ".")
    s = re.sub(r"(\d)_(\d)", r"\1x\2", s)
    s = re.sub(r"(III|II|I)(\d)", r"\1-\2", s, count=1)
    s = re.sub(
        r"(WSL|WDL|WSG)(VIII|VII|IV|VI|IX|V)(\d)",
        r"\1\2-\3",
        s,
        count=1,
        flags=re.I,
    )
    if "/" not in s:
        s = re.sub(r"([YD])(\d+(?:\.\d+)?)-", r"\1/\2-", s, count=1, flags=re.I)
    um_alt = "|".join(re.escape(u) for u in OCTC_UMS)
    s = re.sub(rf"({um_alt})(\d{{1,2}}x\d)", r"\1-\2", s, count=1)
    return s


def parse_wsl_label(raw: str) -> dict | None:
    s = compact_octc_label(raw)
    s = re.sub(
        r"(withhandwheel|withmanualdrive|手动|顶盖.*|头部.*|钟|串并联.*).*$",
        "",
        s,
        flags=re.I,
    )
    m = re.match(
        rf"^(?P<family>WSL|WDL|WSG){OCTC_ROMAN}-"
        r"(?P<current>\d+)"
        r"(?P<conn>[YD])?"
        r"[/-]"
        r"(?P<um>\d+(?:\.\d+)?)"
        r"-"
        r"(?P<tail>.+)$",
        s,
        re.I,
    )
    if not m:
        return None
    tail = m.group("tail")
    pairs = re.findall(r"(\d+)\s*x\s*(\d+)", tail, re.I)
    if not pairs:
        return None
    contact = f"{int(pairs[-1][0])}x{int(pairs[-1][1])}"
    size_m = re.search(r"\(([A-E])\)\s*$|([A-E])\s*$", tail, re.I)
    size = ""
    if size_m:
        size = (size_m.group(1) or size_m.group(2) or "").upper()
    um = float(m.group("um"))
    family = m.group("family").upper()
    series = m.group("series").upper()
    conn = (m.group("conn") or "").upper()
    current = int(m.group("current"))
    list_key = f"{family}{series}-{current}{conn}/{um:g}-{contact}{size}"
    return {
        "family": family,
        "phases": series,
        "currentA": current,
        "connection": conn,
        "umKv": um,
        "selectorSize": size,
        "tapCode": contact,
        "listKey": list_key,
    }


def has_cma7(vals: list) -> bool:
    """True when col E looks like a priced CMA7 set — number is not stored."""
    if len(vals) <= 4:
        return False
    cma7 = vals[4]
    return (
        isinstance(cma7, (int, float))
        and not isinstance(cma7, bool)
        and cma7 > 0
    )


def main() -> None:
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else ONEDRIVE
    if not src.exists():
        sys.exit(f"missing 2025 list xlsx: {src}")

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    if "WSL(WDL)" not in wb.sheetnames:
        sys.exit("sheet WSL(WDL) missing")
    ws = wb["WSL(WDL)"]
    keys: list[str] = []
    octc: list[dict] = []
    seen: set[str] = set()
    for row in ws.iter_rows(max_col=8):
        vals = [c.value for c in row]
        label = next(
            (v.strip() for v in vals if isinstance(v, str) and v.strip()),
            None,
        )
        if not label or not has_cma7(vals):
            continue
        parsed = parse_wsl_label(label)
        if not parsed:
            continue
        key = parsed["listKey"]
        if key in seen:
            continue
        seen.add(key)
        keys.append(key)
        octc.append(parsed)
    wb.close()

    payload = {
        "source": "QS/a. Base Price List 2025.xlsx (WSL/WDL keys only, no RMB)",
        "generatedOn": date.today().isoformat(),
        "keyCount": len(keys),
        "keys": keys,
        "octc": octc,
    }
    text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    if "listRmb" in text or re.search(r"\blistRmb\b", text):
        raise SystemExit("refusing to write listRmb into listIndex")
    OUT.write_text(text, encoding="utf-8")
    print(f"wrote {OUT} keys={len(keys)}")


if __name__ == "__main__":
    main()
